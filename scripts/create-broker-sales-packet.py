from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "BROKER_SALES_PACKET"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "GGuard_Widget_Sales_Broker_Handoff_Packet.docx"

NAVY = "0B2545"; TEAL = "0F766E"; BLUE = "2E74B5"; PALE = "E8EEF5"; LIGHT = "F4F6F9"; GRAY = "5B6573"; RED = "9B1C1C"; GOLD = "7A5A00"

def font(run, size=11, bold=False, color="1F2937", italic=False):
    run.font.name = "Calibri"; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size); run.bold = bold; run.italic = italic; run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd")) or OxmlElement("w:shd"); shd.set(qn("w:fill"), fill)
    if shd.getparent() is None: tcPr.append(shd)

def margins(cell, top=100, start=120, bottom=100, end=120):
    tcPr = cell._tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if tcMar.getparent() is None: tcPr.append(tcMar)
    for side, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")
        if node.getparent() is None: tcMar.append(node)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement("w:tblHeader"); tblHeader.set(qn("w:val"), "true"); trPr.append(tblHeader)

def table_geometry(table, widths):
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr; tblW=tblPr.find(qn("w:tblW")) or OxmlElement("w:tblW"); tblW.set(qn("w:w"),str(sum(widths)));tblW.set(qn("w:type"),"dxa")
    if tblW.getparent() is None: tblPr.append(tblW)
    tblInd=tblPr.find(qn("w:tblInd")) or OxmlElement("w:tblInd");tblInd.set(qn("w:w"),"120");tblInd.set(qn("w:type"),"dxa")
    if tblInd.getparent() is None: tblPr.append(tblInd)
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col=OxmlElement("w:gridCol");col.set(qn("w:w"),str(width));grid.append(col)
    for row in table.rows:
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width/1440); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
            tcW=cell._tc.get_or_add_tcPr().find(qn("w:tcW"));tcW.set(qn("w:w"),str(width));tcW.set(qn("w:type"),"dxa")

def para(doc, text="", size=11, bold=False, color="1F2937", after=6, before=0, align=None, italic=False, keep=False):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.25;p.paragraph_format.keep_with_next=keep
    if align is not None:p.alignment=align
    font(p.add_run(text),size,bold,color,italic);return p

def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f"Heading {level}");p.paragraph_format.keep_with_next=True
    p.add_run(text);return p

def bullet(doc, text):
    p=doc.add_paragraph(style="List Bullet");p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.25;p.paragraph_format.left_indent=Inches(.375);p.paragraph_format.first_line_indent=Inches(-.188);p.add_run(text);return p

def number(doc, text):
    p=doc.add_paragraph(style="List Number");p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.25;p.paragraph_format.left_indent=Inches(.375);p.paragraph_format.first_line_indent=Inches(-.188);p.add_run(text);return p

def callout(doc, label, text, color=TEAL):
    t=doc.add_table(1,1);table_geometry(t,[9360]);c=t.cell(0,0);shade(c,LIGHT);p=c.paragraphs[0];font(p.add_run(label+" "),11,True,color);font(p.add_run(text),11,False,"1F2937");para(doc,"",after=2)

def page_break(doc): doc.add_page_break()

doc=Document(); sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1);sec.header_distance=sec.footer_distance=Inches(.492)
styles=doc.styles
normal=styles["Normal"];normal.font.name="Calibri";normal.font.size=Pt(11);normal.font.color.rgb=RGBColor.from_string("1F2937");normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.25
for name,size,color,before,after in (("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,"1F4D78",10,5)):
    st=styles[name];st.font.name="Calibri";st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(color);st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(header.add_run("GGuard Diagnostics | Authorized Broker Packet | Confidential"),8.5,False,GRAY)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(footer.add_run("Proprietary and confidential - approved sales use only"),8,False,GRAY)

# Cover
para(doc,"AUTHORIZED SALES BROKER ENABLEMENT",10,True,TEAL,4,20)
para(doc,"GGuard Pre-Dispatch",31,True,NAVY,8)
para(doc,"Widget Sales and Contractor Onboarding Handoff Packet",15,False,GRAY,24)
t=doc.add_table(4,2);table_geometry(t,[2700,6660])
for i,(a,b) in enumerate((("Prepared for","Authorized GGuard sales broker"),("Broker portal","https://ggaurdai.com/broker"),("Sales motion","Qualify, demo, onboard, close"),("Document status","Approved controlled-pilot terms"))):
    shade(t.cell(i,0),PALE);font(t.cell(i,0).paragraphs[0].add_run(a),10,True,NAVY);font(t.cell(i,1).paragraphs[0].add_run(b),10.5,False,"1F2937")
para(doc,"",after=12)
callout(doc,"Before you sell:","Sign the broker agreement, activate your broker account, and use only the claims, prices, and materials in this packet.",GOLD)
heading(doc,"Your job",1)
para(doc,"Find qualified garage-door contractors, show the sample workflow, create the deal in the broker portal, send the one-time onboarding link, guide the owner through widget setup, and move the account to paid activation.")
heading(doc,"Your boundaries",1)
bullet(doc,"You are an independent sales channel, not a technician, diagnostician, legal representative, or company administrator.")
bullet(doc,"You cannot access customer media, private service requests, passwords, production credentials, or source code.")
bullet(doc,"You may not promise diagnostic certainty, savings, integrations, compliance, exclusivity, or roadmap features.")

page_break(doc)
heading(doc,"1. What You Are Selling",1)
para(doc,"GGuard Pre-Dispatch is a branded intake widget and hosted upload flow for garage-door service companies. It lets a customer describe a problem and safely provide photos or a short video before a truck rolls. The contractor receives a structured request for dispatcher and technician review.")
callout(doc,"Approved one-sentence pitch:","GGuard Pre-Dispatch helps garage-door companies see what is waiting before dispatch by collecting structured customer details and visual evidence through a branded website widget.")
heading(doc,"Customer value",2)
bullet(doc,"Cleaner information before dispatch.")
bullet(doc,"A branded upload experience on the contractor website or through a hosted link.")
bullet(doc,"Private, tenant-scoped requests and configurable evidence retention.")
bullet(doc,"Preliminary AI-assisted organization of visible evidence, with technician verification required.")
bullet(doc,"Email and secure inbox delivery; signed webhooks are available for generic integration.")
heading(doc,"What it is not",2)
bullet(doc,"It is not a substitute for an on-site inspection or qualified technician.")
bullet(doc,"It does not guarantee a diagnosis, repair, conversion rate, savings, safety outcome, or revenue result.")
bullet(doc,"Planned vendor-specific CRM, SMS, Zapier/Make, and other integrations must not be sold as currently included.")

heading(doc,"2. Ideal Customer",1)
t=doc.add_table(1,3);hdr=t.rows[0];set_repeat_header(hdr)
for c,v in zip(hdr.cells,("Strong fit","Qualifying evidence","Disqualifier")):shade(c,PALE);font(c.paragraphs[0].add_run(v),10,True,NAVY)
rows=[("Residential garage-door service company","Meaningful inbound call volume and a website or willingness to use a hosted link","No owner/operations sponsor"),("Dispatcher wants better pre-visit evidence","Frequent vague calls, repeat questions, avoidable first visits, or difficult scheduling","Expects automatic final diagnosis"),("Owner values branded customer experience","Can install a snippet or authorize a website manager","Refuses consent/privacy requirements"),("Team can review incoming requests","Named inbox owner and test-case commitment","No capacity to act on submissions")]
for row in rows:
    cells=t.add_row().cells
    for c,v in zip(cells,row):font(c.paragraphs[0].add_run(v),9.5);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;margins(c)
table_geometry(t,[2700,3660,3000])

page_break(doc)
heading(doc,"3. Discovery and Qualification",1)
para(doc,"Do not open with a feature dump. Learn the current intake and dispatch process, quantify the friction, then show the relevant workflow.")
heading(doc,"Required discovery questions",2)
for q in ["How do customers currently show you what is wrong before the visit?","How often does the dispatcher receive too little information to prepare the right technician or conversation?","Who reviews customer photos or videos today, and where are they stored?","What happens when a customer cannot clearly describe the door, opener, spring area, or damage?","What website platform do you use, and who can add a small script?","Who owns the buying decision, implementation, and incoming-request inbox?","What result would make a 30-day trial worthwhile?","Are there privacy, insurance, franchise, or CRM approval requirements?"]: bullet(doc,q)
heading(doc,"Qualification standard",2)
callout(doc,"Advance the deal only when:","A decision-maker confirms the problem, a real implementation owner exists, the company can provide an approved HTTPS domain or use the hosted link, and the team will complete a test submission.")
heading(doc,"Pipeline stages",2)
t=doc.add_table(1,3);set_repeat_header(t.rows[0])
for c,v in zip(t.rows[0].cells,("Stage","Broker action","Exit criterion")):shade(c,PALE);font(c.paragraphs[0].add_run(v),10,True,NAVY)
for row in [("Invited","Create deal and send one-time claim link","Prospect receives link"),("Claimed","Confirm owner access and settings","Prospect reaches widget settings"),("Demo","Run approved sample-only demo","Problem and workflow fit confirmed"),("Proposal","Send approved commercial terms","Decision process and date agreed"),("Checkout","Prospect starts Stripe activation","Checkout link issued"),("Won","Verify paid activation","Stripe confirms subscription"),("Lost","Record honest reason","No active follow-up")]:
    cells=t.add_row().cells
    for c,v in zip(cells,row):font(c.paragraphs[0].add_run(v),9.5);margins(c)
table_geometry(t,[1500,3900,3960])

page_break(doc)
heading(doc,"4. Approved Demo Flow",1)
for s in ["Open /pre-dispatch and state the customer problem in one sentence.","Open /pre-dispatch/demo and clearly say that it is sample data; it stores and uploads nothing.","Show the customer intake view: description, photos/video, consent, and safety guidance.","Show the sample preliminary description and repeat that technician verification is required.","Show the request/email/integration views without claiming vendor-specific integrations.","Show the real widget settings only in an approved demo account; never expose another customer tenant.","Close by agreeing on website owner, allowed domain, notification inbox, test submission, and decision date."]: number(doc,s)
heading(doc,"Demo language",2)
t=doc.add_table(1,2);set_repeat_header(t.rows[0])
for c,v in zip(t.rows[0].cells,("Say","Do not say")):shade(c,PALE);font(c.paragraphs[0].add_run(v),10,True,NAVY)
for row in [("Preliminary information for dispatch review","The AI diagnoses the problem"),("A qualified technician verifies findings","This replaces an inspection"),("Private, tenant-scoped workflow","Guaranteed compliance or security"),("Generic signed webhooks are available","It integrates with every CRM"),("Designed to improve intake quality","It will eliminate truck rolls or guarantee savings")]:
    cells=t.add_row().cells
    for c,v in zip(cells,row):font(c.paragraphs[0].add_run(v),10);margins(c)
table_geometry(t,[4680,4680])

page_break(doc)
heading(doc,"5. Commercial Terms and Commission",1)
callout(doc,"CONTROLLED PILOT TERMS:","These are the approved standard terms. Any exception requires written owner approval before it is offered.",RED)
t=doc.add_table(7,2);table_geometry(t,[3300,6060])
fields=[("Approved monthly price","$149/month or $1,490/year"),("Setup or onboarding fee","Waived for standard setup"),("Approved trial","30 days; payment only through approved Stripe checkout"),("Maximum broker discount","No additional discount without written owner approval"),("Broker commission","20% of net collected eligible subscription revenue"),("Commission duration","Customer's first 12 paid months"),("Payout rule","Monthly after a 30-day refund/chargeback hold; signed agreement controls")]
for i,(a,b) in enumerate(fields):shade(t.cell(i,0),PALE);font(t.cell(i,0).paragraphs[0].add_run(a),10,True,NAVY);font(t.cell(i,1).paragraphs[0].add_run(b),10.5)
heading(doc,"Commission operating rule",2)
bullet(doc,"The portal calculates pending commission from the attributed Stripe checkout using the broker rate configured by the owner.")
bullet(doc,"Pending does not mean payable. The owner approves payouts after verifying collected revenue, refunds, chargebacks, taxes, agreement terms, and any clawback period.")
bullet(doc,"The broker may not change attribution, create duplicate deals, self-refer, or promise a discount outside the approved rate card.")

heading(doc,"6. Broker Portal Procedure",1)
for s in ["Sign in at /broker with the authorized broker account.","Create the prospect using the legal/operating company name, unique slug, decision-maker email, phone, and exact HTTPS website URL.","Copy the one-time onboarding link immediately. The raw token is not stored and cannot be displayed again.","Send the link only to the named decision-maker. Their signed-in email must match the invitation.","After they claim the account, guide them through allowed origin, branding, notification email, test submission, and Stripe activation.","Update the deal stage after each meaningful event and record the lost reason honestly.","Review commission status in the broker portal; payment occurs only after admin approval under the signed agreement."]: number(doc,s)

page_break(doc)
heading(doc,"7. Objection Handling",1)
t=doc.add_table(1,2);set_repeat_header(t.rows[0])
for c,v in zip(t.rows[0].cells,("Objection","Approved response")):shade(c,PALE);font(c.paragraphs[0].add_run(v),10,True,NAVY)
responses=[("We already ask for photos.","That is a useful starting point. GGuard adds a consistent branded intake, explicit consent, private tenant-scoped handling, and a structured request before dispatch."),("Will it diagnose the door?","No. It organizes customer information and may provide preliminary visible-evidence observations. A qualified technician must verify the issue."),("Will it work with our website?","The lightweight launcher is designed for common sites and includes a hosted-link option. Installation must be tested on the company's exact allowed HTTPS origin."),("Does it integrate with our CRM?","Email and a focused inbox are available now, along with generic signed webhooks. Do not assume a native vendor integration unless the current approved product sheet lists it."),("What happens to customer media?","The workflow uses private storage intent, tenant scoping, configured retention, and controlled access. Final privacy/security review must follow the customer's requirements."),("Can you guarantee ROI?","No. We can define a trial success measure around intake completeness, dispatcher usefulness, activation, and repeat use, then evaluate actual results.")]
for row in responses:
    cells=t.add_row().cells
    for c,v in zip(cells,row):font(c.paragraphs[0].add_run(v),9.5);margins(c,120,140,120,140)
table_geometry(t,[2800,6560])

page_break(doc)
heading(doc,"8. Onboarding and Close Checklist",1)
heading(doc,"Before sending the claim link",2)
for x in ["Broker agreement signed and account authorized","Decision-maker email verified","Company name and unique slug confirmed","Exact HTTPS website URL confirmed or hosted-link decision made","Price/trial/commission terms approved","No conflicting broker or duplicate deal"]:bullet(doc,"[ ] "+x)
heading(doc,"Before checkout",2)
for x in ["Prospect claimed owner account","Allowed website origin saved","Notification email verified","Widget or hosted link tested","One complete test request submitted","Privacy/terms reviewed by prospect","Decision-maker accepts approved commercial terms"]:bullet(doc,"[ ] "+x)
heading(doc,"After checkout",2)
for x in ["Stripe shows active/trialing subscription","Broker deal automatically shows Won","Commission shows Pending","Customer receives support/implementation contact","Admin verifies payment and approves commission when eligible","Broker records next customer-success check-in"]:bullet(doc,"[ ] "+x)

heading(doc,"9. Security, Confidentiality, and IP",1)
bullet(doc,"Use only approved collateral. Do not copy source code, prompts, schemas, internal reports, customer lists, or nonpublic operating material.")
bullet(doc,"Do not register domains, social handles, ads, or marks containing GGuard without written approval.")
bullet(doc,"Do not share prospect links, customer information, screenshots, or credentials outside the approved deal.")
bullet(doc,"Report suspected copying, impersonation, security incidents, or improper access immediately.")
bullet(doc,"At termination, stop using all GGuard materials and complete required return/deletion certification.")
callout(doc,"Portal access is conditional:","The owner may disable a broker immediately. Access does not grant ownership, source rights, customer rights, or authority to bind GGuard beyond a signed agreement.",RED)

page_break(doc)
heading(doc,"10. First 30 Days",1)
t=doc.add_table(1,4);set_repeat_header(t.rows[0])
for c,v in zip(t.rows[0].cells,("Week","Activity","Minimum output","Review")):shade(c,PALE);font(c.paragraphs[0].add_run(v),9.5,True,NAVY)
rows=[("1","Learn product, sign terms, build named list","25 qualified contractors","Owner certifies demo"),("2","Discovery outreach and demos","10 conversations, 3 demos","Message and objections"),("3","Create qualified deals and onboard","3 claim links, 2 claimed","Pipeline hygiene"),("4","Advance proposals and close","2 proposals, 1 paid target","Conversion and next plan")]
for row in rows:
    cells=t.add_row().cells
    for c,v in zip(cells,row):font(c.paragraphs[0].add_run(v),9.2);margins(c)
table_geometry(t,[900,3300,3000,2160])
heading(doc,"Weekly broker review",2)
bullet(doc,"Deals by stage and next action/date.")
bullet(doc,"Qualified versus disqualified prospects and reasons.")
bullet(doc,"Claims used, onboarding failures, test submissions, and checkout status.")
bullet(doc,"Unapproved questions, objections, requested features, and compliance concerns.")
bullet(doc,"Pending, approved, paid, refunded, or void commissions.")

heading(doc,"11. Sign-off",1)
para(doc,"Broker acknowledgement: I will use only approved claims and materials, protect confidential information and prospect data, avoid unauthorized commitments, and follow the signed broker agreement and portal rules.")
for label in ["Broker name","Broker signature","Date","Owner name","Owner signature","Effective date"]:
    para(doc,f"{label}: " + "_"*55,10.5,False,"1F2937",10)
para(doc,"This packet is an operating guide, not the broker agreement. If they conflict, the signed agreement controls.",9.5,False,GRAY,6,8,italic=True)

doc.save(DOCX)
print(DOCX)
